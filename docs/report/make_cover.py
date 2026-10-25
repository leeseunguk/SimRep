# /// script
# dependencies = ["python-docx"]
# ///
"""Cover page source for build_report.py --mode append (the skill has no title block)."""
import sys
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH

template, out = sys.argv[1], sys.argv[2]
doc = Document(template)
body = doc.element.body
for child in list(body):
    if not child.tag.endswith("sectPr"):
        body.remove(child)
for _ in range(6):
    doc.add_paragraph("")
doc.add_paragraph("SimRep: SIMION과 같은 기능을 더 빠르게 계산하는 네이티브 이온 광학 시뮬레이터", style="Title")
for line in ["SIMION_Rep 최종 기술 보고서", "", "2026년 10월", "", "저장소: SIMION_Rep (커밋 전 작업 트리 기준)"]:
    p = doc.add_paragraph(line)
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
doc.save(out)
print("cover", out)
